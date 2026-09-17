"""Update the 5-slide GISEC deck with the real-export validation and ATHAR-MCP.

Text-only, surgical edits: the deck's layout is full and its design system is good, so
nothing is moved or restyled — only the wording of existing runs changes, which preserves
every font, colour and position. Run from the repo root:

    python hack_files/update_deck.py
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation

SRC = Path("hack_files/ATHAR_GISEC_Hackathon-2.pptx")
BAK = Path("hack_files/ATHAR_GISEC_Hackathon-2.original.pptx")
OUT = Path("hack_files/ATHAR_GISEC_Hackathon-final.pptx")


def para_runs(shape):
    return [r for p in shape.text_frame.paragraphs for r in p.runs]


def set_run(shape, index: int, text: str) -> None:
    """Replace one run's text, keeping its formatting."""
    runs = para_runs(shape)
    if index < len(runs):
        runs[index].text = text


def find(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"{name} not on slide")


def collapse(shape, text: str) -> None:
    """Put all text in run 0 and blank the rest (for runs split mid-phrase)."""
    runs = para_runs(shape)
    if not runs:
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def main() -> None:
    if not BAK.exists():
        shutil.copy2(SRC, BAK)
        print(f"backed up original → {BAK.name}")

    prs = Presentation(str(SRC))
    s1, _s2, s3, s4, s5 = prs.slides

    # ---------------------------------------------------------------- slide 1
    # Lead line picks up the two new capabilities without changing the layout.
    set_run(
        find(s1, "Text 4"),
        0,
        "Unified AWS · Azure · GCP identity governance — blast-radius-ranked risk, "
        "proven on real cloud exports, with a verifiable on-chain audit trail.",
    )

    # ---------------------------------------------------------------- slide 3
    # The agent card now names MCP, which is how other tools reach ATHAR.
    set_run(find(s3, "Text 53"), 0, "LLM agents · ATHAR-MCP")
    set_run(find(s3, "Text 54"), 0, "Investigate · plan fix · MCP tools any AI can call")
    # The fence card states the MCP boundary explicitly.
    set_run(
        find(s3, "Text 74"),
        0,
        "Severity, score and action are deterministic. Over ATHAR-MCP any AI can read and "
        "propose — never approve or apply. API roles enforce it.",
    )

    # ---------------------------------------------------------------- slide 4
    set_run(
        find(s4, "Text 3"),
        0,
        "Every claim here was run, not asserted — including a pass over a real AWS IAM export.",
    )
    # The R7 limitation is now a result: it fires on real data.
    lim = find(s4, "Text 35")
    runs = para_runs(lim)
    if len(runs) >= 2:
        runs[0].text = "R7 peer-outlier now fires on real data  "
        runs[1].text = "36 hits on a real AWS export."
    else:
        collapse(lim, "R7 peer-outlier now fires on real data — 36 hits on a real AWS export.")

    # ---------------------------------------------------------------- slide 5
    set_run(
        find(s5, "Text 3"),
        0,
        "Synthetic demo estate at month 12 · fictional customer “Nahar Digital Authority” "
        "· plus a real AWS IAM export",
    )
    # Fourth headline stat becomes the real-world result the jury asked for.
    collapse(find(s5, "Text 17"), "318")
    set_run(find(s5, "Text 18"), 0, "findings on a real AWS export")
    set_run(find(s5, "Text 19"), 0, "122 real principals · R7 fired 36×")

    # MCP appears in what a governance team actually gets.
    set_run(find(s5, "Text 24"), 0, "Drafted fixes — via dashboard or MCP — human-approved with SoD")

    # Limits and next steps move on now that real data has been run.
    limits = find(s5, "Text 27")
    lr = para_runs(limits)
    if len(lr) >= 2:
        lr[0].text = "LIMITS TODAY  "
        lr[1].text = "Single-host writer key · simulated apply · mapping coverage on real AWS"
    collapse(find(s5, "Text 36"), "Live government estate")

    prs.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
